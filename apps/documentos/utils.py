import os

def extraer_texto_archivo(ruta_archivo):
    """
    Extrae el contenido textual de archivos PDF, DOCX, XLSX, TXT y CSV.
    """
    if not os.path.exists(ruta_archivo):
        return ""

    ext = os.path.splitext(ruta_archivo)[1].lower()
    texto = ""

    try:
        if ext == '.pdf':
            import pypdf
            reader = pypdf.PdfReader(ruta_archivo)
            for i, page in enumerate(reader.pages):
                contenido = page.extract_text()
                if contenido:
                    texto += f"--- Página {i+1} ---\n" + contenido + "\n\n"

        elif ext in ['.docx', '.doc']:
            import docx
            doc = docx.Document(ruta_archivo)
            parrafos = [p.text for p in doc.paragraphs if p.text.strip()]
            texto += "\n".join(parrafos)

            # Extraer también de tablas si existen
            for table in doc.tables:
                for row in table.rows:
                    celdas = [c.text.strip() for c in row.cells if c.text.strip()]
                    if celdas:
                        texto += "\n| " + " | ".join(celdas) + " |"

        elif ext in ['.xlsx', '.xls']:
            import openpyxl
            wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                texto += f"\n--- Hoja: {sheetname} ---\n"
                for row in ws.iter_rows(values_only=True):
                    celdas = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if celdas:
                        texto += " | ".join(celdas) + "\n"

        elif ext in ['.txt', '.csv', '.md']:
            with open(ruta_archivo, 'r', encoding='utf-8', errors='ignore') as f:
                texto = f.read()

    except Exception as e:
        texto = f"[Aviso: No se pudo extraer automáticamente el texto de {ext}: {str(e)}]"

    return texto.strip()
